#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
docbuild —— 项目文档的排版工具箱（统一风格版）
================================================
三份文档（学习笔记 / 完全教程 / 逐行手册）共用的一身「排版衣服」：
统一的字体、标题样式、代码底色、行距、页边距和页码。

用法：调用 h1 / h2 / h3 / para / code / bullet / note / qa 往文档写字，
最后 save() 保存。想换风格，改最上面那组常量即可。
"""

from datetime import datetime

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

AUTHOR = '刘梓涵'                     # 文档属性里的作者

# ---------- 统一配色与字体（想换风格就改这里） ----------
BODY_FONT = 'Microsoft YaHei'        # 正文字体（Windows 自带，显示稳定）
BODY_SIZE = 10.5                     # 正文字号（首行缩进按它折算）
CODE_FONT = 'Consolas'               # 代码字体（等宽）；中文部分自动用正文字体兜底
C_H1 = RGBColor(0x1B, 0x3A, 0x5C)    # 一级标题：深藏青
C_H2 = RGBColor(0x2A, 0x5A, 0x84)    # 二级标题：蓝
C_H3 = RGBColor(0x3D, 0x5A, 0x73)    # 三级标题：灰蓝
C_GRAY = RGBColor(0x6B, 0x6B, 0x6B)  # 备注用的灰色
CODE_BG = 'F4F6F8'                   # 代码块底色（很淡的灰蓝）

doc = Document()


def _get_rfonts(rpr):
    """在「字符属性」里找到字体设置袋（没有就创建一个）。"""
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    return rfonts


def _run_cjk(run, name=BODY_FONT):
    """给一段文字指定「中文字体」（eastAsia）——否则中文会用系统随机的兜底字体。"""
    rpr = run._element.get_or_add_rPr()
    _get_rfonts(rpr).set(qn('w:eastAsia'), name)


def _style_font(style, name):
    """把某个样式的西文和中文（eastAsia）字体都设成同一个。"""
    style.font.name = name
    _get_rfonts(style.element.get_or_add_rPr()).set(qn('w:eastAsia'), name)


def _indent_first_line(paragraph, chars=2):
    """首行缩进 N 个字符（中文段落的老规矩）。

    只给「正文段落」用：Word 里要显示成「首行缩进 2 字符」，靠的是
    w:firstLineChars 这个以「字符」为单位的属性（python-docx 没有接口，
    得自己塞进 w:ind）；同时写一个磅值的 firstLine 兜底，
    免得 WPS / LibreOffice 这类不认 Chars 的软件当没看见。
    """
    pf = paragraph.paragraph_format
    pf.first_line_indent = Pt(BODY_SIZE * chars)
    ind = paragraph._p.get_or_add_pPr().find(qn('w:ind'))
    if ind is not None:
        ind.set(qn('w:firstLineChars'), str(int(chars * 100)))


def _shade(paragraph, fill):
    """给整个段落刷一层底色（代码块的浅灰底就靠它）。"""
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    pPr.append(shd)


def _setup_page():
    """A4 纸 + 舒适页边距 + 页脚页码。"""
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.3)
    sec.left_margin = sec.right_margin = Cm(2.4)

    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    f2 = OxmlElement('w:instrText'); f2.set(qn('xml:space'), 'preserve'); f2.text = 'PAGE'
    f3 = OxmlElement('w:fldChar'); f3.set(qn('w:fldCharType'), 'end')
    run._r.append(f1); run._r.append(f2); run._r.append(f3)
    run.font.size = Pt(9)
    run.font.color.rgb = C_GRAY


def _setup_styles():
    """全局样式：正文、各级标题、列表的字体与间距。"""
    normal = doc.styles['Normal']
    _style_font(normal, BODY_FONT)
    normal.font.size = Pt(BODY_SIZE)
    normal.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    normal.paragraph_format.line_spacing = 1.4
    # 中文排版习惯：段与段之间不留空行（靠首行缩进分段），所以段后间距为 0
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)

    for name, size, color in [('Heading 1', 19, C_H1),
                              ('Heading 2', 14, C_H2),
                              ('Heading 3', 12, C_H3)]:
        st = doc.styles[name]
        _style_font(st, BODY_FONT)
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.italic = False
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(12 if name == 'Heading 2' else 10)
        st.paragraph_format.space_after = Pt(5)
        st.paragraph_format.line_spacing = 1.2

    # 封面大标题用的 Title 样式
    st = doc.styles['Title']
    _style_font(st, BODY_FONT)
    st.font.size = Pt(30)
    st.font.bold = True
    st.font.italic = False
    st.font.color.rgb = C_H1

    lb = doc.styles['List Bullet']
    lb.paragraph_format.line_spacing = 1.35
    lb.paragraph_format.space_after = Pt(0)


_setup_page()
_setup_styles()


def cover(title, subtitle, meta):
    """封面页：大标题 + 副标题 + 说明（全部居中）。"""
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(34)
    r.font.color.rgb = C_H1
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(subtitle)
    r2.font.size = Pt(14)
    r2.font.color.rgb = C_H2
    for _ in range(5):
        doc.add_paragraph()
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p3.add_run(meta)
    r3.font.size = Pt(10.5)
    r3.font.color.rgb = C_GRAY


def h1(text, new_page=True):
    """一级标题（章）。默认另起一页。"""
    if new_page:
        doc.add_page_break()
    doc.add_heading(text, level=1)


def h2(text):
    """二级标题（节）。"""
    doc.add_heading(text, level=2)


def h3(text):
    """三级标题（小节）。"""
    doc.add_heading(text, level=3)


def para(text, indent=True):
    """正文段落：首行缩进 2 字符，段与段之间不空行。

    少数不适合缩进的（比如目录式的条目）传 indent=False。
    """
    p = doc.add_paragraph(text)
    if indent:
        _indent_first_line(p)
    return p


def bullet(text):
    """圆点列表项。"""
    doc.add_paragraph(text, style='List Bullet')


def code(text):
    """代码 / 命令 / 输出块：等宽字体 + 浅灰蓝底色 + 左侧缩进。"""
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = CODE_FONT
    r.font.size = Pt(9)
    _run_cjk(r, BODY_FONT)          # 代码里的中文注释用正文字体显示
    pf = p.paragraph_format
    # 代码块上下各留一点点缝就够（底色本身就是分隔符），不占整行
    pf.space_before = Pt(3)
    pf.space_after = Pt(3)
    pf.line_spacing = 1.25
    pf.left_indent = Cm(0.35)
    _shade(p, CODE_BG)
    return p


def note(text):
    """灰色斜体小提示（小坑、冷知识）。"""
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(10)
    r.font.color.rgb = C_GRAY
    p.paragraph_format.left_indent = Cm(0.35)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    return p


def qa(q, a):
    """问答对（FAQ 用）：问题加粗顶格，答案首行缩进。"""
    p = doc.add_paragraph()
    r = p.add_run('Q：' + q)
    r.bold = True
    p.paragraph_format.space_before = Pt(4)
    ans = doc.add_paragraph('A：' + a)
    _indent_first_line(ans)
    ans.paragraph_format.space_after = Pt(4)


def image(path, caption=None, width_cm=14.5, height_cm=None):
    """插入一张配图（居中，可选灰色小字说明）。
    width_cm / height_cm 二选一控制大小（竖长图用 height_cm）。"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if height_cm is not None:
        p.add_run().add_picture(path, height=Cm(height_cm))
    else:
        p.add_run().add_picture(path, width=Cm(width_cm))
    if caption:
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(caption)
        r.italic = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = C_GRAY
    return p


def save(path, title=None, created=None):
    """保存文档（自动建目录）。

    不设的话，python-docx 会在文档属性里留下「作者：python-docx」
    和「说明：generated by python-docx」，创建时间还是它模板里写死的
    2013-12-23。这些东西在资源管理器里右键「属性→详细信息」就能看到，
    所以这里一并写成正常的值。created 传 'YYYY-MM-DD HH:MM' 字符串。
    """
    import os
    cp = doc.core_properties
    cp.author = AUTHOR
    cp.last_modified_by = AUTHOR
    cp.comments = ''          # dc:description，python-docx 默认塞了「generated by python-docx」
    cp.category = ''
    cp.title = title or os.path.splitext(os.path.basename(path))[0]
    cp.revision = 1
    if created:
        t = datetime.strptime(created, '%Y-%m-%d %H:%M')
        cp.created = t
        cp.modified = t

    d = os.path.dirname(path)
    if d:
        os.makedirs(d, exist_ok=True)
    doc.save(path)
    print('已生成', path)
